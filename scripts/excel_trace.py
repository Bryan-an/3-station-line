# scripts/excel_trace.py
"""CLI: run one replication and export the event trace to .xlsx.

Used to produce the deliverable 2.2 spreadsheet manually.

Example:
    python scripts/excel_trace.py --k1 3 --k2 3 --seed 12345 --duration 480 --out trace.xlsx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

# Allow running as `python scripts/excel_trace.py` from project root
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from src.simulation import SimulationConfig, run_simulation


def export_trace(result, config: SimulationConfig, out: Path) -> None:
    wb = Workbook()

    # --- Sheet 1: Trace ---
    ws = wb.active
    ws.title = "Trace"
    headers = [
        "t (min)", "Tipo", "Pieza", "U usado", "T generado (min)",
        "Estado E1", "Estado E2", "Estado E3",
        "Pieza E1", "Pieza E2", "Pieza E3",
        "B1 cont.", "B2 cont.", "WIP",
        "Piezas terminadas", "Próx. pieza ID", "Comentario",
    ]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2C3E50")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for ev in result.event_log:
        ws.append([
            round(ev["t"], 4),
            ev["event_type"],
            ev["piece_id"],
            round(ev["u"], 6) if ev["u"] is not None else "",
            round(ev["duration"], 4) if ev["duration"] is not None else "",
            ev["e1_state"], ev["e2_state"], ev["e3_state"],
            ev["e1_piece"], ev["e2_piece"], ev["e3_piece"],
            ev["b1_count"], ev["b2_count"], ev["wip"],
            ev["pieces_completed"], ev["next_piece_id"], ev["note"],
        ])
    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(10, min(max_len + 2, 40))

    # --- Sheet 2: Métricas ---
    ms = wb.create_sheet("Métricas")
    ms.append(["Métrica", "Valor"])
    for cell in ms[1]:
        cell.font = header_font
        cell.fill = header_fill
    rows = [
        ("Piezas completadas", result.pieces_completed),
        ("Throughput (piezas/h)", round(result.throughput, 4)),
        ("Tiempo de flujo promedio (min)",
         round(sum(result.flow_times) / len(result.flow_times), 4) if result.flow_times else 0.0),
        ("WIP promedio", round(result.avg_wip, 4)),
        ("% Bloqueo E1", round(result.pct_blocked_e1 * 100, 3)),
        ("% Bloqueo E2", round(result.pct_blocked_e2 * 100, 3)),
        ("% Hambre E2", round(result.pct_starved_e2 * 100, 3)),
        ("% Hambre E3", round(result.pct_starved_e3 * 100, 3)),
        ("% Ocupado E1", round(result.busy_e1 * 100, 3)),
        ("% Ocupado E2", round(result.busy_e2 * 100, 3)),
        ("% Ocupado E3", round(result.busy_e3 * 100, 3)),
    ]
    for r in rows:
        ms.append(r)
    ms.column_dimensions["A"].width = 35
    ms.column_dimensions["B"].width = 15

    # --- Sheet 3: Inputs ---
    ins = wb.create_sheet("Inputs")
    ins.append(["Parámetro", "Valor"])
    for cell in ins[1]:
        cell.font = header_font
        cell.fill = header_fill
    for k, v in {
        "K1": config.k1,
        "K2": config.k2,
        "Tasa λ (1/min)": config.rate,
        "Media exponencial (min)": 1.0 / config.rate,
        "Duración (min)": config.duration,
        "Semilla GLC": config.seed,
        "GLC a": 1_664_525,
        "GLC c": 1_013_904_223,
        "GLC m": 2 ** 32,
    }.items():
        ins.append([k, v])
    ins.column_dimensions["A"].width = 25
    ins.column_dimensions["B"].width = 20

    wb.save(out)


def main() -> int:
    parser = argparse.ArgumentParser(description="Export 3-station-line trace to .xlsx")
    parser.add_argument("--k1", type=int, required=True)
    parser.add_argument("--k2", type=int, required=True)
    parser.add_argument("--seed", type=int, default=12_345)
    parser.add_argument("--duration", type=float, default=480.0)
    parser.add_argument("--rate", type=float, default=0.25)
    parser.add_argument("--out", type=Path, default=Path("trace.xlsx"))
    args = parser.parse_args()

    cfg = SimulationConfig(
        k1=args.k1, k2=args.k2, rate=args.rate, duration=args.duration, seed=args.seed
    )
    result = run_simulation(cfg)
    export_trace(result, cfg, args.out)
    print(f"Wrote {args.out} — {result.pieces_completed} pieces, "
          f"{result.throughput:.2f} piezas/h, {len(result.event_log)} events")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
